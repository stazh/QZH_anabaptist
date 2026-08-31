#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extrahiert Header-Infos aus TEI/XML in Excel.
- Fix: <note>-Inhalte innerhalb von <head> werden ignoriert.
- Fix: Deep-Text Extraktion für <head> (behält <persName> bei).
"""

import copy
from pathlib import Path
import argparse
import re
from lxml import etree
import openpyxl
from openpyxl.styles import Font, Alignment

# Fester Namespace für TEI
NS = {'tei': 'http://www.tei-c.org/ns/1.0'}

FIELDS = [
    ("index", "Nr."),                                  # A
    ("head", "Titel/Head"),                            # B
    ("idno", "Vorlage/Idno (Überlieferungsträger)"),   # C
    ("origDate", "Datum (origDate)"),                  # D
    ("ueberlieferung", "Überlieferung"),               # E
    ("transcript", "Transkript (Personen)"),           # F
    ("tagging", "Tagging (Personen)"),                 # G
    ("publisher", "Publisher"),                        # H
    ("series_idno", "Serien-ID"),                      # I
    ("source_title", "Serientitel"),                   # J
]

def get_clean_head_text(head_elements):
    """
    Extrahiert Text aus <head>, ignoriert aber <note>-Elemente.
    Behält Texte in <persName> etc. bei.
    """
    if head_elements is not None and len(head_elements) > 0:
        # Wir arbeiten auf einer Kopie, um das Original-XML nicht zu verändern
        head_copy = copy.deepcopy(head_elements[0])
        
        # Alle <note> Elemente innerhalb des Heads entfernen
        notes = head_copy.xpath(".//tei:note", namespaces=NS) or head_copy.xpath(".//note")
        for note in notes:
            parent = note.getparent()
            if parent is not None:
                # Falls Text nach der Note kommt (Tail), muss dieser erhalten bleiben
                if note.tail:
                    previous = note.getprevious()
                    if previous is not None:
                        previous.tail = (previous.tail or "") + note.tail
                    else:
                        parent.text = (parent.text or "") + note.tail
                parent.remove(note)
        
        # Jetzt den verbleibenden Text extrahieren
        text = head_copy.xpath("string(.)")
        return text.strip() if text else ""
    return ""

def safe_get_text(element_list):
    """Standard Deep-Text Extraktion für andere Felder."""
    if element_list is not None and len(element_list) > 0:
        text = element_list[0].xpath("string(.)")
        return text.strip() if text else ""
    return ""

def extract_header_info_from_div(div):
    """Extrahiert Metadaten aus einem <div type='document'>-Block."""

    def xp(path):
        res = div.xpath(path, namespaces=NS)
        if not res:
            res = div.xpath(path.replace("tei:", ""))
        return res

    # 1. Head (Titel) - Spezialreinigung für <note>
    head = get_clean_head_text(xp(".//tei:head"))

    # 2. Überlieferungsträger & Überlieferung
    idno_str = ""
    ueberlieferung_str = ""
    carrier_divs = xp(".//tei:div[@type='original' or @type='vorlage' or @type='Abschrift' or @type='Edition']")
    
    base_parts = []
    ueberlieferung_parts = []

    for s_div in carrier_divs:
        idno_el = s_div.xpath(".//tei:idno", namespaces=NS) or s_div.xpath(".//idno")
        if idno_el:
            raw_text = safe_get_text(idno_el)
        else:
            raw_text = " ".join(s_div.itertext()).strip()

        if not raw_text:
            continue

        parens = re.findall(r"\(([^()]*)\)", raw_text)
        if parens:
            ueberlieferung_parts.extend([p.strip() for p in parens if p.strip()])
            no_paren = re.sub(r"\s*\([^()]*\)", "", raw_text)
        else:
            no_paren = raw_text

        no_paren = no_paren.strip().rstrip(",;")
        if no_paren:
            base_parts.append(no_paren)

    idno_str = "; ".join(base_parts)
    ueberlieferung_str = "; ".join(ueberlieferung_parts)

    # 3. Datum (origDate)
    origDate = ""
    date_divs = xp(".//tei:div[@type='source']")
    if date_divs:
        src = date_divs[0]
        date_el = src.xpath(".//tei:origDate", namespaces=NS) or src.xpath(".//origDate")
        if date_el:
            origDate = date_el[0].get("when") or date_el[0].xpath("string(.)").strip()
        else:
            origDate = src.get("date", "")

    # 4. RespStmt
    transcribers, taggers = [], []
    for r in xp(".//tei:respStmt"):
        resp = r.find(".//tei:resp", namespaces=NS) or r.find(".//resp")
        pers = r.find(".//tei:persName", namespaces=NS) or r.find(".//persName")
        if resp is not None and pers is not None:
            name = pers.xpath("string(.)").strip()
            role = resp.get("key")
            if role == "transcript": transcribers.append(name)
            elif role == "tagging": taggers.append(name)

    return {
        "head": head,
        "idno": idno_str,
        "origDate": origDate.strip(),
        "ueberlieferung": ueberlieferung_str,
        "transcript": ", ".join(transcribers),
        "tagging": ", ".join(taggers),
        "publisher": safe_get_text(xp(".//tei:publisher")),
        "series_idno": safe_get_text(xp(".//tei:idno")),
        "source_title": safe_get_text(xp(".//tei:title")),
    }

def main():
    ap = argparse.ArgumentParser(description="TEI Metadaten-Extraktor")
    ap.add_argument("input_xml", type=Path)
    ap.add_argument("output_xlsx", type=Path)
    args = ap.parse_args()

    if not args.input_xml.exists():
        print(f"Datei nicht gefunden: {args.input_xml}")
        return

    parser = etree.XMLParser(remove_blank_text=True, recover=True)
    tree = etree.parse(str(args.input_xml), parser)
    root = tree.getroot()
    doc_divs = root.xpath("//tei:div[@type='document']", namespaces=NS) or root.xpath("//div[@type='document']")

    print(f"Verarbeite {len(doc_divs)} Dokumente...")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEI Metadaten"
    ws.append([label for _, label in FIELDS])
    for cell in ws[1]: cell.font = Font(bold=True)

    for idx, div in enumerate(doc_divs, start=1):
        try:
            info = extract_header_info_from_div(div)
            row = [idx] + [info.get(key, "") for key, _ in FIELDS if key != "index"]
            ws.append(row)
        except Exception as e:
            print(f"Fehler in Dok {idx}: {e}")
            ws.append([idx, "FEHLER"])

    # Spaltenbreite & Alignment
    for col_let, width in [('C', 60), ('E', 40)]:
        ws.column_dimensions[col_let].width = width
        for cell in ws[col_let]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    wb.save(str(args.output_xlsx))
    print(f"Erfolgreich gespeichert: {args.output_xlsx}")

if __name__ == "__main__":
    main()