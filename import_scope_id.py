import os
import re
from openpyxl import load_workbook
from lxml import etree

# ==========================================
# KONFIGURATION
# ==========================================
EXCEL_FILE = r'C:\Temp\Repositories\QZH_anabaptist\Scope_IDs_bestehende_VEs.xlsx'
XML_FOLDER = r'C:\Temp\Repositories\QZH_anabaptist\input'
BASE_URL = 'https://suche.staatsarchiv.djiktzh.ch/detail.aspx?ID='

def get_comparison_string(text):
    """
    Erstellt einen temporären String für den Abgleich.
    Ignoriert runde Klammern und deren Inhalt.
    """
    if not text:
        return ""
    # Ignoriert alles von ( bis )
    no_brackets = re.sub(r'\(.*?\)', '', str(text))
    # Entfernt überschüssige Leerzeichen für einen sauberen Abgleich
    return " ".join(no_brackets.split())

def main():
    print("Lese Excel-Datei ein...")
    wb = load_workbook(EXCEL_FILE, data_only=True)
    ws = wb.active
    
    mapping = {}
    for row in ws.iter_rows(min_row=1, min_col=2, max_col=3, values_only=True):
        col_b_val, col_c_val = row
        if col_b_val is not None and col_c_val is not None:
            comparison_b = get_comparison_string(col_b_val)
            # Wir speichern jetzt ein Tuple: (Die ID für den Link, Der Originaltext aus Spalte B für den Log)
            mapping[comparison_b] = (str(col_c_val).strip(), str(col_b_val).strip())

    print(f"{len(mapping)} Einträge aus Excel geladen.")
    print("-" * 40)

    # Set zum Speichern aller erfolgreich gefundenen Einträge
    matched_keys = set()

    for filename in os.listdir(XML_FOLDER):
        if not filename.endswith('.xml'):
            continue
            
        filepath = os.path.join(XML_FOLDER, filename)
        modified = False
        
        try:
            tree = etree.parse(filepath)
            root = tree.getroot()
            
            xpath_query = '//*[local-name()="msIdentifier"]//*[local-name()="idno"]'
            
            for idno in root.xpath(xpath_query):
                original_text = idno.text
                if not original_text:
                    continue
                    
                comparison_idno = get_comparison_string(original_text)
                
                if comparison_idno in mapping:
                    # Werte aus dem Dictionary entpacken
                    id_value, original_excel_text = mapping[comparison_idno]
                    
                    # Den Schlüssel als "erfolgreich gefunden" markieren
                    matched_keys.add(comparison_idno)
                    
                    new_source_url = f"{BASE_URL}{id_value}"
                    idno.set('source', new_source_url)
                    modified = True
                    print(f"[{filename}] Match gefunden für: '{comparison_idno}'. Attribut hinzugefügt.")
            
            if modified:
                tree.write(filepath, encoding='utf-8', xml_declaration=True, pretty_print=True)
                
        except Exception as e:
            print(f"Fehler beim Verarbeiten von {filename}: {e}")

    print("-" * 40)
    
    # ==========================================
    # AUSWERTUNG DER NICHT GEFUNDENEN EINTRÄGE
    # ==========================================
    # Differenz bilden: Alle Excel-Keys minus die gefundenen Keys
    unmatched_keys = set(mapping.keys()) - matched_keys
    
    if unmatched_keys:
        print(f"WARNUNG: Es wurden {len(unmatched_keys)} Einträge aus der Excel-Datei in keinem XML gefunden:")
        for key in unmatched_keys:
            # Den originalen Excel-Text aus dem Mapping holen
            _, original_excel_text = mapping[key]
            print(f" - {original_excel_text}")
    else:
        print("Perfekt! Alle Einträge aus der Excel-Datei wurden in den XML-Dateien gefunden.")

    print("-" * 40)
    print("Vorgang abgeschlossen!")

if __name__ == "__main__":
    main()